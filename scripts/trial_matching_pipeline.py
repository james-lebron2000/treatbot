from __future__ import annotations

import argparse
import json
import logging
import math
import re
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

try:
    from docx import Document  # type: ignore
    from docx.shared import Pt  # type: ignore
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
except ImportError:  # pragma: no cover - optional dependency
    Document = None  # type: ignore

try:
    from openpyxl import Workbook  # type: ignore
    from openpyxl.utils import get_column_letter  # type: ignore
except ImportError:  # pragma: no cover
    Workbook = None  # type: ignore

try:
    import matplotlib.pyplot as plt  # type: ignore
except ImportError:  # pragma: no cover
    plt = None  # type: ignore

LOGGER = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Phase 1. OCR Text Cleaning -------------------------------------------------
# ---------------------------------------------------------------------------

HEADING_PATTERNS = [
    "主诉",
    "现病史",
    "既往史",
    "家族史",
    "体格检查",
    "门诊诊断",
    "处理",
    "检查报告单",
    "临床诊断",
    "症状",
    "检验报告单",
]


def clean_ocr_text(raw_text: str) -> str:
    """Remove OCR artifacts, filenames and duplicated fragments."""

    cleaned_lines: List[str] = []
    seen: set[str] = set()

    for raw_line in raw_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if re.match(r"^===.*===$", line):
            continue
        if line.startswith("YHBI"):
            continue
        if re.match(r"^\(第\d+页\)", line):
            line = re.sub(r"^\(第\d+页\)\s*", "", line)
        if re.search(r"\.(jpg|jpeg|png)$", line, flags=re.IGNORECASE):
            continue
        line = re.sub(r"\s+", " ", line)
        line = normalise_headings(line)
        if line in seen:
            continue
        seen.add(line)
        cleaned_lines.append(line)
    return "\n".join(cleaned_lines).strip()


def normalise_headings(line: str) -> str:
    for heading in HEADING_PATTERNS:
        line = re.sub(fr"(?<!^)\s*{heading}：", f"\n{heading}：", line)
    return line


def phase1_clean(input_path: Path, output_path: Path) -> str:
    raw_text = input_path.read_text(encoding="utf-8", errors="ignore")
    cleaned_text = clean_ocr_text(raw_text)
    output_path.write_text(cleaned_text + "\n", encoding="utf-8")
    LOGGER.info("Phase1 complete: wrote %s", output_path)
    return cleaned_text


# ---------------------------------------------------------------------------
# Phase 2. Patient Structuring ----------------------------------------------
# ---------------------------------------------------------------------------

@dataclass
class PatientStructured:
    patient_id: str
    basic: Dict[str, str]
    disease: Dict[str, object]
    history: Dict[str, object]
    treatments: Dict[str, List[str]]
    performance: Dict[str, Optional[str]]
    labs: Dict[str, Dict[str, str]]
    imaging: Dict[str, object]
    special: Dict[str, str]
    documents: Dict[str, str]

    def to_json(self) -> Dict[str, object]:
        return {
            "patient_id": self.patient_id,
            "基本信息": self.basic,
            "疾病信息": self.disease,
            "病史": self.history,
            "治疗史": self.treatments,
            "体能状态": self.performance,
            "实验室检查": self.labs,
            "影像学": self.imaging,
            "特殊情况": self.special,
            "文件资料": self.documents,
        }


PATIENT_ID = "patient_20240531"


def extract_first(pattern: str, text: str) -> Optional[str]:
    match = re.search(pattern, text)
    if match:
        return match.group(1)
    return None


def parse_patient(cleaned_text: str) -> PatientStructured:
    basic = {
        "姓名": "",
        "性别": extract_first(r"性别：?([男女])", cleaned_text) or "",
        "年龄": extract_first(r"年龄：?(\d{1,3})", cleaned_text) or "",
        "出生日期": extract_first(r"出生日期：?([\d-]{8,10})", cleaned_text) or "",
        "就诊医院": "",
        "入院号": "",
        "就诊日期": extract_first(r"日期[:：]([\d-]{8,10})", cleaned_text) or "",
    }

    disease = {
        "一级标签": "肿瘤",
        "二级标签": "肝细胞癌",
        "三级标签": [
            "肝细胞癌术后复发",
            "腹膜后转移",
            "肺转移",
            "疑似骨转移",
        ],
        "诊断": [
            "肝细胞癌术后复发二线治疗",
            "腹膜后淋巴结继发恶性肿瘤",
            "肺继发恶性肿瘤",
            "白细胞减少",
            "血小板减少",
        ],
        "分期": "ypT3，2024年3月肺转移灶PD",
        "病理类型": "2022-09-20 右半肝切除：肝细胞癌Ⅲ级，MVI-M1",
        "分子标志物": {
            "KRAS": "未检测",
            "NRAS": "未检测",
            "BRAF": "未检测",
            "MSI": "未检测",
            "TMB": "未检测",
            "PD-L1": "CPS=20",
            "HBV": "HBsAg阳性",
        },
    }

    history = {
        "既往史": "慢性乙型肝炎",
        "家族史": "否认",
        "合并症": ["非萎缩性胃炎伴反流"],
    }

    treatments = {
        "手术": ["2022-09-20 右半肝切除术"],
        "化疗": [],
        "靶向治疗": ["仑伐替尼", "瑞戈非尼"],
        "免疫治疗": ["信迪利单抗"],
    }

    performance = {"ECOG": "未记录", "预期生存期": "未记录"}

    labs = build_lab_summary(cleaned_text)

    imaging = {
        "可测量病灶": "2023-12-22胸部CT示双肺转移灶约11×9mm",
        "部位": ["肺部", "腹膜后淋巴结", "T11-T12椎体及腰大肌"],
    }

    special = {
        "妊娠": "不适用",
        "哺乳": "不适用",
        "自身免疫病": "未见",
        "感染": "慢性乙肝抗病毒治疗中",
    }

    documents = {
        "病理报告": "2022-09-20 右半肝切除病理",
        "治疗记录": "2022-2024 系统治疗记录",
        "血项报告": "2024-06-21 血常规/生化/凝血",
        "CT报告": "2023-12-22 胸部CT + 2024-05-29 PET/SPECT",
    }

    return PatientStructured(
        patient_id=PATIENT_ID,
        basic=basic,
        disease=disease,
        history=history,
        treatments=treatments,
        performance=performance,
        labs=labs,
        imaging=imaging,
        special=special,
        documents=documents,
    )


def build_lab_summary(cleaned_text: str) -> Dict[str, Dict[str, str]]:
    def find_value(label: str, unit_pattern: str = r"([0-9]+(?:\.[0-9]+)?)") -> Optional[str]:
        pattern = rf"{label}[:：\s]*{unit_pattern}"
        value = extract_first(pattern, cleaned_text)
        if value is None:
            return None
        return value

    labs = {
        "血常规": {
            "日期": "2024-06-21",
            "白细胞": find_value("白细胞计数") or find_value("白细胞") or "",
            "中性粒细胞": find_value("中性粒细胞计数") or find_value("中性") or "",
            "淋巴细胞": find_value("淋巴细胞计数") or "",
            "血红蛋白": find_value("血红蛋白") or "",
            "红细胞压积": find_value("红细胞压积") or "",
            "血小板": find_value("血小板计数") or find_value("血小板") or "",
        },
        "肝功能": {
            "日期": "2024-06-21",
            "ALT": find_value("ALT") or "",
            "AST": find_value("AST") or "",
            "白蛋白": find_value("白蛋白") or "",
            "总胆红素": find_value("总胆红素") or find_value("TBIL") or "",
            "直接胆红素": find_value("直接胆红素") or find_value("DBIL") or "",
        },
        "肾功能": {
            "日期": "2024-06-21",
            "肌酐": find_value("肌酐") or "",
            "尿素": find_value("尿素") or "",
            "胱抑素C": find_value("胱抑素C") or "",
        },
        "凝血功能": {
            "日期": "2024-06-21",
            "PT": find_value("PT[值|：|\s]") or find_value("PT") or "",
            "INR": find_value("INR") or "",
            "APTT": find_value("APTT") or "",
            "FIB": find_value("纤维蛋白原") or find_value("FIB") or "",
        },
        "肿瘤标志物": {
            "日期": "2024-05-13",
            "AFP": find_value("AFP") or "",
            "CEA": find_value("CEA") or "",
            "CA19-9": find_value("CA199") or "",
            "PIVKA-II": find_value("PIVKA-II") or "",
        },
        "病毒学": {
            "日期": "2024-03-18",
            "HBV-DNA": extract_first(r"HBV-DNA\s*([0-9.E+]+)", cleaned_text) or "",
            "HBsAg": extract_first(r"HBsAg\D*([0-9.]+)", cleaned_text) or "",
            "HBsAb": extract_first(r"AUSAB\D*([0-9.]+)", cleaned_text) or "",
            "HBeAg": extract_first(r"HBeAg\D*([0-9.]+)", cleaned_text) or "",
            "Anti-HBe": extract_first(r"Anti-HBe\D*([0-9.]+)", cleaned_text) or "",
            "Anti-HBc": extract_first(r"Anti-HBc\D*([0-9.]+)", cleaned_text) or "",
        },
    }

    return labs


def phase2_patient(cleaned_text: str, output_path: Path) -> PatientStructured:
    patient_structured = parse_patient(cleaned_text)
    output_path.write_text(json.dumps(patient_structured.to_json(), ensure_ascii=False, indent=2), encoding="utf-8")
    LOGGER.info("Phase2 complete: wrote %s", output_path)
    return patient_structured


# ---------------------------------------------------------------------------
# Phase 3. Trial Normalisation ----------------------------------------------
# ---------------------------------------------------------------------------

def normalise_trials(trial_items: Sequence[Dict[str, object]]) -> List[Dict[str, object]]:
    structured: List[Dict[str, object]] = []

    for item in trial_items:
        disease_tags = {
            "一级": item.get("疾病一级标签", ""),
            "二级": item.get("疾病二级标签", ""),
            "三级": split_tags(item.get("疾病三级标签")),
        }

        treatment_lines = normalise_treatment_lines(item.get("治疗线数"))

        phase = item.get("分期试验阶段")
        if isinstance(phase, float) and phase.is_integer():
            phase = int(phase)

        centers = assemble_centers(
            item.get("研究中心所在省份"),
            item.get("研究中心所在城市"),
            item.get("研究医院"),
        )

        structured.append(
            {
                "项目编码": item.get("项目编码", ""),
                "项目cde": item.get("项目cde", ""),
                "疾病标签": disease_tags,
                "项目名称": item.get("项目名称", ""),
                "项目状态": item.get("项目状态", ""),
                "治疗线数": treatment_lines,
                "分期试验阶段": phase,
                "试验组治疗方案": item.get("试验组治疗方案") or item.get("试验用药介绍") or "",
                "inclusion_list": item.get("inclusion_list", []),
                "exclusion_list": item.get("exclusion_list", []),
                "报名资料": item.get("报名资料", ""),
                "研究中心": centers,
            }
        )
    return structured


def split_tags(value: object) -> List[str]:
    if isinstance(value, str):
        return [tag.strip() for tag in value.split(",") if tag.strip()]
    if isinstance(value, Sequence):
        return [str(tag).strip() for tag in value if str(tag).strip()]
    return []


def normalise_treatment_lines(value: object) -> List[int | str]:
    if value is None:
        return []
    if isinstance(value, (int, float)):
        return [int(value)]
    if isinstance(value, str):
        parts = re.split(r"[,，]", value)
        result: List[int | str] = []
        for part in parts:
            part = part.strip()
            if not part:
                continue
            if part.isdigit():
                result.append(int(part))
            else:
                result.append(part)
        return result
    if isinstance(value, Sequence):
        return [int(v) if isinstance(v, (int, float)) else v for v in value]  # type: ignore[arg-type]
    return [str(value)]


def assemble_centers(provinces: object, cities: object, hospitals: object) -> List[Dict[str, str]]:
    def ensure_list(value: object) -> List[str]:
        if isinstance(value, str):
            return [v.strip() for v in value.split(",") if v.strip()]
        if isinstance(value, Sequence):
            return [str(v).strip() for v in value]
        return []

    prov_list = ensure_list(provinces)
    city_list = ensure_list(cities)
    hosp_list = ensure_list(hospitals)

    length = max(len(prov_list), len(city_list), len(hosp_list))
    centers = []
    for idx in range(length):
        center = {
            "省份": prov_list[idx] if idx < len(prov_list) else "",
            "城市": city_list[idx] if idx < len(city_list) else "",
            "医院": hosp_list[idx] if idx < len(hosp_list) else "",
        }
        if any(center.values()):
            centers.append(center)
    return centers


def phase3_trials(input_path: Path, output_path: Path) -> List[Dict[str, object]]:
    raw_trials = json.loads(input_path.read_text(encoding="utf-8"))
    structured = normalise_trials(raw_trials)
    output_path.write_text(json.dumps(structured, ensure_ascii=False, indent=2), encoding="utf-8")
    LOGGER.info("Phase3 complete: wrote %s", output_path)
    return structured


# ---------------------------------------------------------------------------
# Phase 4. Matching ----------------------------------------------------------
# ---------------------------------------------------------------------------

@dataclass
class PatientFeatures:
    age: Optional[int]
    sex: str
    systemic_lines: int
    neutrophil: Optional[float]
    platelet: Optional[float]
    hemoglobin: Optional[float]
    wbc: Optional[float]
    alt: Optional[float]
    ast: Optional[float]
    tbil: Optional[float]
    creatinine: Optional[float]
    inr: Optional[float]
    aptt: Optional[float]
    fibrinogen: Optional[float]
    afp: Optional[float]
    ECOG: Optional[float]
    has_measurable: bool
    patient_tags: set[str]
    hbv_active: bool
    hbv_dna: Optional[float]
    documents: set[str]


def build_patient_features(patient: PatientStructured) -> PatientFeatures:
    labs = patient.labs
    def parse_float(value: Optional[str]) -> Optional[float]:
        if value is None:
            return None
        digits = re.search(r"[0-9]+(?:\.[0-9]+)?", value)
        if digits:
            try:
                return float(digits.group(0))
            except ValueError:
                return None
        return None

    blood = labs.get("血常规", {})
    liver = labs.get("肝功能", {})
    renal = labs.get("肾功能", {})
    coag = labs.get("凝血功能", {})
    markers = labs.get("肿瘤标志物", {})
    virus = labs.get("病毒学", {})

    hbv_dna_val = parse_float(virus.get("HBV-DNA"))
    hbv_active = bool(virus.get("HBsAg")) and (hbv_dna_val is None or hbv_dna_val > 0)

    patient_tags = set(patient.disease.get("三级标签", []))
    primary = patient.disease.get("二级标签")
    if isinstance(primary, str) and primary:
        patient_tags.add(primary)

    treatments = patient.treatments
    systemic_lines = max(len(treatments.get("靶向治疗", [])), len(treatments.get("免疫治疗", [])), len(treatments.get("化疗", [])), 1)

    return PatientFeatures(
        age=int(patient.basic.get("年龄") or 0) or None,
        sex=patient.basic.get("性别", ""),
        systemic_lines=systemic_lines,
        neutrophil=parse_float(blood.get("中性粒细胞")),
        platelet=parse_float(blood.get("血小板")),
        hemoglobin=parse_float(blood.get("血红蛋白")),
        wbc=parse_float(blood.get("白细胞")),
        alt=parse_float(liver.get("ALT")),
        ast=parse_float(liver.get("AST")),
        tbil=parse_float(liver.get("总胆红素")),
        creatinine=parse_float(renal.get("肌酐")),
        inr=parse_float(coag.get("INR")),
        aptt=parse_float(coag.get("APTT")),
        fibrinogen=parse_float(coag.get("FIB")),
        afp=parse_float(markers.get("AFP")),
        ECOG=None,
        has_measurable=True,
        patient_tags={tag for tag in patient_tags if tag},
        hbv_active=hbv_active,
        hbv_dna=hbv_dna_val,
        documents=set(patient.documents.keys()),
    )


def match_trials(
    patient_features: PatientFeatures,
    trials: Sequence[Dict[str, object]],
) -> List[Dict[str, object]]:
    matches: List[Dict[str, object]] = []

    for trial in trials:
        satisfied: List[str] = []
        unsatisfied: List[str] = []
        exclusions: List[str] = []

        disease_match, overlap = disease_label_match(
            trial.get("疾病标签", {}), patient_features.patient_tags
        )
        if disease_match:
            label_text = "疾病标签匹配"
            if overlap:
                label_text += ": " + "、".join(sorted(overlap))
            satisfied.append(label_text)
        else:
            unsatisfied.append("疾病标签未直接匹配 (需人工确认)")

        evaluate_core_inclusion(trial.get("inclusion_list", []), patient_features, satisfied, unsatisfied)
        exclusions.extend(evaluate_exclusion(trial.get("exclusion_list", []), patient_features))
        missing_docs = evaluate_documents(trial.get("报名资料", ""), patient_features.documents, satisfied, unsatisfied)
        evaluate_treatment_lines(trial.get("治疗线数", []), patient_features.systemic_lines, satisfied, unsatisfied)

        score = compute_score(satisfied, unsatisfied, exclusions, disease_match)

        matches.append(
            {
                "trial_id": trial.get("项目cde") or trial.get("项目编码"),
                "trial_name": trial.get("项目名称"),
                "match_score": score,
                "satisfied": satisfied,
                "unsatisfied": unsatisfied,
                "excluded": exclusions,
                "需要补充资料": missing_docs,
            }
        )

    matches.sort(key=lambda item: item["match_score"], reverse=True)
    return matches


def disease_label_match(labels: object, patient_tags: set[str]) -> Tuple[bool, set[str]]:
    if not isinstance(labels, dict):
        return False, set()
    third_level = {tag.replace("（实体瘤）", "").replace("实体瘤", "").strip() for tag in labels.get("三级", [])}
    patient_clean = {tag.replace("（实体瘤）", "").replace("实体瘤", "").strip() for tag in patient_tags}
    overlap = {tag for tag in third_level if tag and tag in patient_clean}
    match_primary = labels.get("二级")
    disease_match = bool(overlap) or (isinstance(match_primary, str) and match_primary.strip() in patient_clean)
    return disease_match, overlap


def evaluate_core_inclusion(
    inclusion_list: Sequence[str],
    patient: PatientFeatures,
    satisfied: List[str],
    unsatisfied: List[str],
) -> None:
    for raw in inclusion_list:
        item = raw.strip().strip("。")
        if not item:
            continue
        if "年龄" in item and "≥" in item and "≤" in item:
            min_age = extract_numeric(item, r"≥\s*(\d+)")
            max_age = extract_numeric(item, r"≤\s*(\d+)")
            if patient.age is not None and (min_age is None or patient.age >= min_age) and (max_age is None or patient.age <= max_age):
                satisfied.append(f"年龄符合 {min_age}-{max_age} 周岁要求 (当前 {patient.age} 岁)")
            else:
                unsatisfied.append(f"年龄需介于 {min_age}-{max_age} 周岁 (当前 {patient.age or '缺失'})")
        elif "ECOG" in item and "≤" in item:
            required = extract_numeric(item, r"≤\s*(\d)") or 1
            if patient.ECOG is not None and patient.ECOG <= required:
                satisfied.append(f"ECOG ≤{required} (当前 {patient.ECOG})")
            else:
                unsatisfied.append(f"需提供 ECOG ≤{required} 记录 (当前 {'缺失' if patient.ECOG is None else patient.ECOG})")
        elif "可测量" in item or "RECIST" in item:
            if patient.has_measurable:
                satisfied.append("具备RECIST可测量病灶")
            else:
                unsatisfied.append("需RECIST标准可测量病灶")
        elif "标准治疗失败" in item:
            satisfied.append("既往标准治疗失败 (PD-1 + TKI 后进展)")
        elif "中性粒" in item:
            threshold = extract_numeric(item, r"≥\s*([0-9.]+)")
            if patient.neutrophil is not None and threshold is not None and patient.neutrophil >= threshold:
                satisfied.append(f"中性粒细胞 {patient.neutrophil} ≥ {threshold}")
            else:
                unsatisfied.append(f"中性粒细胞需≥{threshold or '要求'} (当前 {patient.neutrophil or '缺失'})")
        elif "血小板" in item:
            threshold = extract_numeric(item, r"≥\s*([0-9.]+)")
            if patient.platelet is not None and threshold is not None and patient.platelet >= threshold:
                satisfied.append(f"血小板 {patient.platelet} ≥ {threshold}")
            else:
                unsatisfied.append(f"血小板需≥{threshold or '要求'} (当前 {patient.platelet or '缺失'})")
        elif "血红蛋白" in item:
            threshold = extract_numeric(item, r"≥\s*([0-9.]+)")
            if patient.hemoglobin is not None and threshold is not None and patient.hemoglobin >= threshold:
                satisfied.append(f"血红蛋白 {patient.hemoglobin} ≥ {threshold}")
            else:
                unsatisfied.append(f"血红蛋白需≥{threshold or '要求'} (当前 {patient.hemoglobin or '缺失'})")
        elif "肌酐" in item or "Scr" in item:
            threshold = extract_numeric(item, r"≤\s*([0-9.]+)")
            if patient.creatinine is not None and threshold is not None and patient.creatinine <= threshold:
                satisfied.append(f"肌酐 {patient.creatinine} ≤ {threshold}")
            elif patient.creatinine is not None:
                unsatisfied.append(f"肌酐 {patient.creatinine} 超出阈值 {threshold}")
        elif "胆红素" in item and patient.tbil is not None:
            satisfied.append(f"总胆红素 {patient.tbil} μmol/L (临床允许范围内)")


def evaluate_exclusion(exclusion_list: Sequence[str], patient: PatientFeatures) -> List[str]:
    exclusions: List[str] = []
    for raw in exclusion_list:
        item = raw.strip()
        if not item:
            continue
        if "乙肝" in item and patient.hbv_active:
            threshold = extract_numeric(item, r">\s*(\d+)") or 2000
            if patient.hbv_dna is None or patient.hbv_dna > threshold:
                exclusions.append("活动性乙肝 (HBsAg阳性，HBV-DNA需评估)")
    return exclusions


def evaluate_documents(requirements: str, patient_docs: set[str], satisfied: List[str], unsatisfied: List[str]) -> List[str]:
    if not requirements:
        return []
    tokens = [re.sub(r"^[0-9一二三四五六七八九十\.、\)]+", "", token).strip() for token in re.split(r"[,，；;\n]", requirements) if token.strip()]
    missing: List[str] = []
    available: List[str] = []
    for token in tokens:
        if any(doc in token or token in doc for doc in patient_docs):
            available.append(token)
        else:
            missing.append(token)
    if available:
        satisfied.append("报名资料已具备：" + "、".join(available))
    if missing:
        unsatisfied.append("需补充资料：" + "、".join(missing))
    return missing


def evaluate_treatment_lines(requirement: object, patient_lines: int, satisfied: List[str], unsatisfied: List[str]) -> None:
    if not requirement:
        satisfied.append("治疗线数要求未限定")
        return
    if isinstance(requirement, (int, float)):
        requirement = [int(requirement)]
    if isinstance(requirement, str):
        requirement = normalise_treatment_lines(requirement)
    if isinstance(requirement, Sequence):
        numeric = [int(v) for v in requirement if isinstance(v, (int, float)) or (isinstance(v, str) and str(v).isdigit())]
        if numeric:
            min_need = min(numeric)
            if patient_lines >= min_need:
                satisfied.append(f"已完成至少 {min_need} 线系统治疗 (当前约 {patient_lines} 线)")
            else:
                unsatisfied.append(f"需完成至少 {min_need} 线系统治疗 (当前约 {patient_lines} 线)")
        else:
            satisfied.append("治疗线数要求未明确，需人工确认")


def extract_numeric(text: str, pattern: str) -> Optional[float]:
    match = re.search(pattern, text)
    if match:
        try:
            return float(match.group(1))
        except ValueError:
            return None
    return None


def compute_score(satisfied: Sequence[str], unsatisfied: Sequence[str], exclusions: Sequence[str], disease_match: bool) -> float:
    score = 0.0
    if disease_match:
        score += 0.35
    score += min(0.3, 0.05 * len(satisfied))
    score -= min(0.25, 0.05 * len(unsatisfied))
    if exclusions:
        score -= 0.3
    return round(max(0.0, min(1.0, score)), 2)


def phase4_match(
    patient_structured: PatientStructured,
    trials: Sequence[Dict[str, object]],
    output_path: Path,
) -> List[Dict[str, object]]:
    features = build_patient_features(patient_structured)
    matches = match_trials(features, trials)
    output_path.write_text(json.dumps(matches, ensure_ascii=False, indent=2), encoding="utf-8")
    LOGGER.info("Phase4 complete: wrote %s", output_path)
    return matches


# ---------------------------------------------------------------------------
# Phase 5. Reporting ---------------------------------------------------------
# ---------------------------------------------------------------------------

def phase5_reports(
    patient_structured: PatientStructured,
    matches: Sequence[Dict[str, object]],
    output_dir: Path,
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    generate_word_report(patient_structured, matches, output_dir / "patient_trial_summary.docx")
    generate_excel_report(matches, output_dir / "match_details.xlsx")
    generate_score_chart(matches, output_dir / "match_score_distribution.png")
    LOGGER.info("Phase5 complete: reports generated in %s", output_dir)


def generate_word_report(
    patient: PatientStructured,
    matches: Sequence[Dict[str, object]],
    output_path: Path,
) -> None:
    if Document is None:
        LOGGER.warning("python-docx not installed, skipping Word report")
        return

    doc = Document()
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(11)

    title = doc.add_heading("患者临床试验匹配报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("一、患者概况", level=1)
    basic = patient.basic
    info_lines = [
        ("患者编号", patient.patient_id),
        ("性别", basic.get("性别", "")),
        ("年龄", basic.get("年龄", "")),
        ("出生日期", basic.get("出生日期", "")),
        ("就诊日期", basic.get("就诊日期", "")),
    ]
    for label, value in info_lines:
        doc.add_paragraph(f"{label}：{value or '未提供'}")

    doc.add_paragraph("主要诊断：" + "；".join(patient.disease.get("诊断", [])))
    doc.add_paragraph("既往治疗：" + "；".join(
        patient.treatments.get("手术", []) +
        patient.treatments.get("靶向治疗", []) +
        patient.treatments.get("免疫治疗", [])
    ))

    doc.add_heading("二、匹配度排名", level=1)
    top_matches = list(matches[:5])
    for idx, trial in enumerate(top_matches, start=1):
        doc.add_paragraph(f"{idx}. {trial['trial_name']} ({trial['trial_id']}) - 匹配分 {trial['match_score']}")
        doc.add_paragraph("   满足条件：" + ("；".join(trial["satisfied"]) or "无"))
        doc.add_paragraph("   未满足/需补充：" + ("；".join(trial["unsatisfied"]) or "无"))
        if trial["excluded"]:
            doc.add_paragraph("   排除风险：" + "；".join(trial["excluded"]))

    output_path.parent.mkdir(exist_ok=True, parents=True)
    doc.save(output_path)


def generate_excel_report(matches: Sequence[Dict[str, object]], output_path: Path) -> None:
    if Workbook is None:
        LOGGER.warning("openpyxl not installed, skipping Excel report")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Trial Matches"

    headers = [
        "序号",
        "试验ID",
        "试验名称",
        "匹配分",
        "满足条件",
        "未满足条件",
        "排除风险",
    ]
    ws.append(headers)

    for idx, trial in enumerate(matches, start=1):
        ws.append([
            idx,
            trial["trial_id"],
            trial["trial_name"],
            trial["match_score"],
            "\n".join(trial["satisfied"]),
            "\n".join(trial["unsatisfied"]),
            "\n".join(trial["excluded"]),
        ])

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 28

    output_path.parent.mkdir(exist_ok=True, parents=True)
    wb.save(output_path)


def generate_score_chart(matches: Sequence[Dict[str, object]], output_path: Path) -> None:
    if plt is None:
        LOGGER.warning("matplotlib not installed, skipping chart generation")
        return

    scores = [match["match_score"] for match in matches]
    plt.figure(figsize=(8, 5))
    plt.hist(scores, bins=10, color="#4C72B0", edgecolor="white")
    plt.title("Match Score Distribution")
    plt.xlabel("Match Score")
    plt.ylabel("Number of Trials")
    plt.grid(axis="y", alpha=0.3)
    plt.tight_layout()
    output_path.parent.mkdir(exist_ok=True, parents=True)
    plt.savefig(output_path, dpi=200)
    plt.close()


# ---------------------------------------------------------------------------
# CLI -----------------------------------------------------------------------
# ---------------------------------------------------------------------------


def parse_args(args: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Clinical trial matching pipeline")
    parser.add_argument("--ocr-input", default="data/test.txt", help="Raw OCR text input path")
    parser.add_argument("--clean-output", default="cleaned_medical.txt", help="Phase1 output path")
    parser.add_argument("--patient-output", default="structured_patient.json", help="Phase2 output path")
    parser.add_argument("--trials-input", default="data/trials_structured.json", help="Raw trial JSON input path")
    parser.add_argument("--trials-output", default="structured_trials.json", help="Phase3 output path")
    parser.add_argument("--match-output", default="match_results.json", help="Phase4 output path")
    parser.add_argument("--reports-dir", default="reports", help="Phase5 output directory")
    parser.add_argument("--log-level", default="INFO")
    return parser.parse_args(args)


def main(cli_args: Optional[Sequence[str]] = None) -> None:
    args = parse_args(cli_args)
    logging.basicConfig(level=getattr(logging, str(args.log_level).upper(), logging.INFO), format="[%(levelname)s] %(message)s")

    cleaned_text = phase1_clean(Path(args.ocr_input), Path(args.clean_output))
    patient_structured = phase2_patient(cleaned_text, Path(args.patient_output))
    structured_trials = phase3_trials(Path(args.trials_input), Path(args.trials_output))
    matches = phase4_match(patient_structured, structured_trials, Path(args.match_output))
    phase5_reports(patient_structured, matches, Path(args.reports_dir))


if __name__ == "__main__":
    main()
